#!/usr/bin/env python3
"""Erzeugt die Schulvorstellung Mathsachs als PowerPoint (16:9)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "Mathsachs-Schulvorstellung.pptx"
LOGO_BI = ROOT / "public/supporters/logo-bi-menschenskinder.png"
LOGO_MEIN = ROOT / "public/supporters/logo-mein-delitzsch.png"

# Tiefes Indigo, hoher Kontrast — gut für Beamer im Klassenraum
NAVY = RGBColor(0x0B, 0x1E, 0x3D)
NAVY_DEEP = RGBColor(0x07, 0x14, 0x2B)
CARD = RGBColor(0x14, 0x32, 0x5C)
CARD_EDGE = RGBColor(0x2A, 0x4A, 0x7A)
GOLD = RGBColor(0xF5, 0xC5, 0x42)
TEAL = RGBColor(0x5E, 0xEA, 0xC4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xD7, 0xE4, 0xF2)
MUTED = RGBColor(0x9B, 0xB4, 0xCC)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _set_run_font(run, *, size: int, bold: bool = False, color: RGBColor = WHITE, name: str = "Calibri") -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    # latin + ea so umlauts stay Calibri in PowerPoint
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _fill_line(shape, fill: RGBColor, line: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)


def add_textbox(slide, left, top, width, height, text, *, size=28, bold=False, color=WHITE, align=PP_ALIGN.LEFT, name="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size=size, bold=bold, color=color, name=name)
    return box


def add_footer(slide, page: int, total: int = 10) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.22), SLIDE_W, Inches(0.28))
    _fill(bar, NAVY_DEEP)
    add_textbox(
        slide,
        Inches(0.45),
        Inches(7.22),
        Inches(8.5),
        Inches(0.28),
        "Mathsachs  ·  Mathe üben nach Lehrplan  ·  v0.1.42",
        size=11,
        color=MUTED,
    )
    add_textbox(
        slide,
        Inches(11.4),
        Inches(7.22),
        Inches(1.5),
        Inches(0.28),
        f"{page} / {total}",
        size=11,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def base_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    _fill(bg, NAVY)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), SLIDE_H)
    _fill(stripe, GOLD)
    return slide


def add_title_block(slide, title: str, kicker: str | None = None) -> None:
    if kicker:
        add_textbox(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.36), kicker, size=14, bold=True, color=GOLD)
        title_top = Inches(0.55)
    else:
        title_top = Inches(0.38)
    add_textbox(slide, Inches(0.55), title_top, Inches(12.2), Inches(0.7), title, size=34, bold=True, color=WHITE)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.28), Inches(2.2), Inches(0.045))
    _fill(rule, GOLD)


def add_bullet_cards(slide, items: list[str], *, top=Inches(1.55), left=Inches(0.55), width=Inches(12.2)) -> None:
    gap = Inches(0.12)
    height = Inches(0.82)
    for i, text in enumerate(items):
        y = top + i * (height + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, y, width, height)
        _fill_line(card, CARD, CARD_EDGE)
        # python-pptx adjustments for corner
        try:
            card.adjustments[0] = 0.12
        except Exception:
            pass
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, left + Inches(0.22), y + Inches(0.24), Inches(0.34), Inches(0.34)
        )
        _fill(marker, GOLD)
        num = slide.shapes.add_textbox(left + Inches(0.22), y + Inches(0.26), Inches(0.34), Inches(0.32))
        p = num.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(i + 1)
        _set_run_font(run, size=13, bold=True, color=NAVY)
        add_textbox(
            slide,
            left + Inches(0.72),
            y + Inches(0.18),
            width - Inches(0.95),
            Inches(0.5),
            text,
            size=22,
            color=WHITE,
        )


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1 Titel ----------------------------------------------------------------
    s = base_slide(prs)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.12), Inches(1.55), Inches(13.213), Inches(4.05))
    _fill(band, CARD)
    add_textbox(s, Inches(0.7), Inches(0.72), Inches(12), Inches(0.4), "Schulvorstellung", size=16, bold=True, color=GOLD)
    add_textbox(s, Inches(0.7), Inches(1.75), Inches(12), Inches(1.0), "Mathsachs", size=60, bold=True, color=WHITE)
    add_textbox(
        s,
        Inches(0.7),
        Inches(2.85),
        Inches(12),
        Inches(0.65),
        "Mathe üben nach Lehrplan",
        size=32,
        color=GOLD,
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(3.55),
        Inches(12),
        Inches(0.45),
        "Gymnasium und Oberschule Sachsen  ·  Hauptschul- und Realschulbildungsgang",
        size=18,
        color=MIST,
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(4.35),
        Inches(12),
        Inches(0.4),
        "Linus und Matthias Ulrich",
        size=18,
        color=WHITE,
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(4.8),
        Inches(12),
        Inches(0.35),
        "Aktuelle Version  v0.1.42",
        size=16,
        color=MUTED,
    )
    add_footer(s, 1)

    # 2 Was ist Mathsachs? ---------------------------------------------------
    s = base_slide(prs)
    add_title_block(s, "Was ist Mathsachs?", "Kurz erklärt")
    lead = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.55), Inches(12.2), Inches(1.15))
    _fill_line(lead, CARD, CARD_EDGE)
    try:
        lead.adjustments[0] = 0.08
    except Exception:
        pass
    add_textbox(
        s,
        Inches(0.8),
        Inches(1.75),
        Inches(11.7),
        Inches(0.85),
        "Mathsachs ist ein Übungsprogramm für Mathe nach dem sächsischen Lehrplan – fürs Gymnasium und die Oberschule.",
        size=22,
        color=MIST,
    )
    add_bullet_cards(
        s,
        [
            "Ihr wählt ein Lehrplan-Thema und übt sofort.",
            "Ihr seht direkt, ob die Lösung stimmt.",
            "Bei Fehlern gibt es den Weg Schritt für Schritt.",
        ],
        top=Inches(2.9),
    )
    add_footer(s, 2)

    # 3 So übst du -----------------------------------------------------------
    s = base_slide(prs)
    add_title_block(s, "So übst du", "Am Bildschirm oder auf Papier")
    add_bullet_cards(
        s,
        [
            "Thema aus dem Lehrplan wählen.",
            "Aufgabe lösen – die Auswertung kommt sofort.",
            "Erklärung anzeigen, wenn etwas nicht klappt.",
            "Oder: Übungsblatt mit Lösungen drucken.",
        ],
    )
    add_footer(s, 3)

    # 4 Punkte und Fortschritt ----------------------------------------------
    s = base_slide(prs)
    add_title_block(s, "Punkte und Fortschritt", "Dein Üben bleibt sichtbar")
    add_bullet_cards(
        s,
        [
            "Im Protokoll siehst du den Stand je Thema und die Gesamtpunktzahl.",
            "Mehrere Personen am selben Gerät: „Wer übt heute?“",
            "Optional: Punkte anonym an die Klasse senden.",
            "Am Schul-PC und auf Tablets im WLAN gilt derselbe Stand.",
        ],
    )
    add_footer(s, 4)

    # 5 Challenge ------------------------------------------------------------
    s = base_slide(prs)
    add_title_block(s, "Challenge", "Gemeinsam ein Ziel")
    add_bullet_cards(
        s,
        [
            "Lehrerin oder Lehrer startet eine Challenge für die Klasse.",
            "Ihr übt festgelegte Themen in einem Zeitraum.",
            "Die Punkte zählen extra für die Challenge.",
            "Online sieht man nur Summen – keine Schülernamen.",
        ],
    )
    add_footer(s, 5)

    # 6 Gemeinsam in der Klasse ----------------------------------------------
    s = base_slide(prs)
    add_title_block(s, "Gemeinsam in der Klasse", "Klassencode – kurz für alle")
    add_bullet_cards(
        s,
        [
            "Die Klasse bekommt einen Klassencode.",
            "Punkte kann man an die Klasse schicken – freiwillig.",
            "Online stehen nur Klassenname und Punktesumme.",
            "Rollen in der App: Schüler, Eltern, Klassenlehrer, Lehrer.",
        ],
    )
    add_footer(s, 6)

    # 7 Datenschutz ----------------------------------------------------------
    s = base_slide(prs)
    add_title_block(s, "Datenschutz in einem Satz", "Kein Name auf dem Klassenserver")
    big = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.6), Inches(12.2), Inches(2.15))
    _fill_line(big, CARD, GOLD)
    try:
        big.adjustments[0] = 0.06
    except Exception:
        pass
    add_textbox(
        s,
        Inches(0.9),
        Inches(1.95),
        Inches(11.5),
        Inches(1.5),
        "Auf dem Klassenserver stehen keine Namen, keine Benutzer-IDs und keine E-Mails – nur der Klassenname und anonyme Punktesummen.",
        size=26,
        bold=True,
        color=WHITE,
    )
    add_bullet_cards(
        s,
        [
            "Der Klassencode bleibt in der Klasse – er ist das Geheimnis.",
            "In der App steht der Hinweis unter Datenschutz.",
        ],
        top=Inches(4.05),
    )
    add_footer(s, 7)

    # 8 Unterstützer ---------------------------------------------------------
    s = base_slide(prs)
    add_title_block(s, "Wer hat uns unterstützt?", "Danke!")
    add_textbox(
        s,
        Inches(0.55),
        Inches(1.5),
        Inches(12.2),
        Inches(0.45),
        "Zwei Unterstützer aus Delitzsch stehen in der Fußzeile der App.",
        size=20,
        color=MIST,
    )

    card_w = Inches(5.85)
    card_h = Inches(3.85)
    lefts = (Inches(0.55), Inches(6.9))
    cards_meta = [
        {
            "title": "Bürgerinitiative Menschenskinder Delitzsch! e.V.",
            "url": "www.bi-menschenskinder-delitzsch.de",
            "logo": LOGO_BI,
            "logo_w": Inches(2.15),
            "logo_h": Inches(2.15),
            "bg": WHITE,
        },
        {
            "title": "Mein Delitzsch",
            "url": "",
            "logo": LOGO_MEIN,
            "logo_w": Inches(4.4),
            "logo_h": Inches(1.69),
            "bg": RGBColor(0x00, 0x00, 0x00),
        },
    ]
    for left, meta in zip(lefts, cards_meta):
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.05), card_w, card_h)
        _fill_line(card, CARD, CARD_EDGE)
        try:
            card.adjustments[0] = 0.06
        except Exception:
            pass
        logo_panel_h = Inches(2.25)
        panel = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.2), Inches(2.22), card_w - Inches(0.4), logo_panel_h
        )
        _fill(panel, meta["bg"])
        try:
            panel.adjustments[0] = 0.08
        except Exception:
            pass
        if meta["logo"].is_file():
            lw, lh = meta["logo_w"], meta["logo_h"]
            lx = left + (card_w - lw) / 2
            ly = Inches(2.22) + (logo_panel_h - lh) / 2
            s.shapes.add_picture(str(meta["logo"]), lx, ly, lw, lh)
        add_textbox(
            s,
            left + Inches(0.25),
            Inches(4.55),
            card_w - Inches(0.5),
            Inches(0.7),
            meta["title"],
            size=16,
            bold=True,
            color=WHITE,
        )
        if meta["url"]:
            add_textbox(
                s,
                left + Inches(0.25),
                Inches(5.25),
                card_w - Inches(0.5),
                Inches(0.4),
                meta["url"],
                size=13,
                color=TEAL,
            )
    add_footer(s, 8)

    # 9 So startest du -------------------------------------------------------
    s = base_slide(prs)
    add_title_block(s, "So startest du", "Nächste Schritte")
    add_bullet_cards(
        s,
        [
            "Mathsachs am Schul-PC öffnen – oder das Tablet im WLAN.",
            "Bei „Wer übt heute?“ deinen Übungsnamen wählen.",
            "Lehrplan-Thema wählen – und losüben.",
            "Zu Hause (Windows): Mathsachs-Setup-x.y.z.exe",
        ],
    )
    add_footer(s, 9)

    # 10 Fragen + Kontakt ----------------------------------------------------
    s = base_slide(prs)
    add_title_block(s, "Fragen?", "Wir hören zu")
    add_textbox(
        s,
        Inches(0.55),
        Inches(1.55),
        Inches(12.2),
        Inches(0.55),
        "Was unklar ist, einfach fragen – hier im Raum oder später per Mail.",
        size=22,
        color=MIST,
    )
    add_bullet_cards(
        s,
        [
            "Kontakt: info@my-smart-home-support.de",
            "Autoren: Linus und Matthias Ulrich",
            "Aktuelle Version: v0.1.42",
        ],
        top=Inches(2.25),
    )
    add_textbox(
        s,
        Inches(0.55),
        Inches(5.15),
        Inches(12.2),
        Inches(0.4),
        "Idee oder Feedback? In der App unter Idee / Feedback – oder dieselbe Mailadresse.",
        size=16,
        color=MUTED,
    )
    add_footer(s, 10)

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"geschrieben: {path}")
